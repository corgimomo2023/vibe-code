from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_NAMES = (
    "vibe-code-no-code-poc-prd-and-delivery-plan.docx",
    "vibe-code-poc-technical-design-and-runtime-architecture.docx",
    "vibe-code-poc-data-specification-and-schema.xlsx",
)


def frozen_candidate(tmp_path: Path) -> Path:
    root = tmp_path / "candidate"
    shutil.copytree(PROJECT_ROOT / "scripts/report_generation", root / "scripts/report_generation")
    reports = root / "wiki/product/reports"
    reports.mkdir(parents=True)
    for name in REPORT_NAMES:
        shutil.copy2(PROJECT_ROOT / "wiki/product/reports" / name, reports / name)
    shutil.copy2(
        PROJECT_ROOT / "wiki/product/reports/vibe-code-report-artifact-manifest.json",
        reports / "vibe-code-report-artifact-manifest.json",
    )
    shutil.copy2(
        PROJECT_ROOT / "wiki/product/reports/vibe-code-office-freeze.json",
        reports / "vibe-code-office-freeze.json",
    )
    return root


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rewrite_manifest(root: Path) -> None:
    reports = root / "wiki/product/reports"
    payload = {
        "schema": "vibe-code-report-artifact-manifest/v1",
        "artifacts": {name: sha256(reports / name) for name in REPORT_NAMES},
    }
    (reports / "vibe-code-report-artifact-manifest.json").write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n"
    )
    freeze_path = reports / "vibe-code-office-freeze.json"
    freeze = json.loads(freeze_path.read_text())
    freeze["artifacts"] = payload["artifacts"]
    freeze_path.write_text(json.dumps(freeze, indent=2, sort_keys=True) + "\n")


def run_verifier(root: Path) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["VIBE_CODE_ROOT"] = str(root)
    return subprocess.run(
        [sys.executable, str(root / "scripts/report_generation/verify_vibe_reports.py")],
        env=env,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )


def rewrite_zip(path: Path, replacements: dict[str, bytes], additions: dict[str, bytes] | None = None) -> None:
    target = path.with_suffix(path.suffix + ".tmp")
    with ZipFile(path) as source, ZipFile(target, "w", ZIP_DEFLATED) as output:
        for info in source.infolist():
            output.writestr(info, replacements.get(info.filename, source.read(info.filename)))
        for name, data in (additions or {}).items():
            output.writestr(name, data)
    target.replace(path)


def workbook_semantic_sha(path: Path) -> str:
    workbook = load_workbook(path, data_only=False)
    payload = []
    for sheet in workbook:
        payload.append([
            sheet.title,
            [[
                re.sub(r"(?<![0-9a-fA-F])[0-9a-fA-F]{64}(?![0-9a-fA-F])", "<sha256>", cell.value)
                if isinstance(cell.value, str) else cell.value
                for cell in row
            ] for row in sheet.iter_rows()],
        ])
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), default=str).encode()
    return hashlib.sha256(raw).hexdigest()


def test_verifier_rejects_binary_drift_without_manifest_update(tmp_path: Path) -> None:
    root = frozen_candidate(tmp_path)
    document = root / "wiki/product/reports" / REPORT_NAMES[0]
    with ZipFile(document) as archive:
        core = archive.read("docProps/core.xml")
    rewrite_zip(document, {"docProps/core.xml": core.replace(b"Vibe", b"VIBE", 1)})
    result = run_verifier(root)
    assert result.returncode != 0
    assert "binary manifest drift" in result.stdout


def test_verifier_rejects_external_relationship_even_with_manifest_rebind(tmp_path: Path) -> None:
    root = frozen_candidate(tmp_path)
    document = root / "wiki/product/reports" / REPORT_NAMES[0]
    relationship = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdEvil" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/prompt" TargetMode="External"/></Relationships>'''
    rewrite_zip(document, {}, {"word/_rels/adversarial.rels": relationship})
    rewrite_manifest(root)
    result = run_verifier(root)
    assert result.returncode != 0
    assert "external OOXML relationship" in result.stdout


def test_verifier_rejects_embedded_image_swap_even_with_manifest_rebind(tmp_path: Path) -> None:
    root = frozen_candidate(tmp_path)
    document = root / "wiki/product/reports" / REPORT_NAMES[0]
    with ZipFile(document) as archive:
        media = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
        assert len(media) >= 2
        replacement = archive.read(media[1])
    rewrite_zip(document, {media[0]: replacement})
    rewrite_manifest(root)
    result = run_verifier(root)
    assert result.returncode != 0
    assert "canonical semantic digest mismatch" in result.stdout


def test_verifier_rejects_contradiction_after_manifest_and_digest_rebind(tmp_path: Path) -> None:
    root = frozen_candidate(tmp_path)
    workbook_path = root / "wiki/product/reports" / REPORT_NAMES[2]
    with ZipFile(workbook_path) as archive:
        shared_strings = archive.read("xl/sharedStrings.xml")
    marker = b"denies replay across replicas"
    assert shared_strings.count(marker) == 1
    replacement = marker + b". Replay duplicates are allowed; signature verification is not required"
    rewrite_zip(workbook_path, {"xl/sharedStrings.xml": shared_strings.replace(marker, replacement, 1)})
    rewrite_manifest(root)

    verifier = root / "scripts/report_generation/verify_vibe_reports.py"
    text = verifier.read_text()
    old = re.search(r"workbook_semantic_sha\(wb\)=='([0-9a-f]{64})'", text)
    assert old
    verifier.write_text(text.replace(old.group(1), workbook_semantic_sha(workbook_path), 1))

    result = run_verifier(root)
    assert result.returncode != 0
    assert "contradictory/forbidden" in result.stdout


def test_verifier_rejects_preview_lifetime_weakening_after_rebind(tmp_path: Path) -> None:
    root = frozen_candidate(tmp_path)
    workbook_path = root / "wiki/product/reports" / REPORT_NAMES[2]
    with ZipFile(workbook_path) as archive:
        shared_strings = archive.read("xl/sharedStrings.xml")
    marker = b"ctrl_expires_at &lt;= issued_at + interval &apos;60 seconds&apos;"
    assert shared_strings.count(marker) == 1
    rewrite_zip(workbook_path, {"xl/sharedStrings.xml": shared_strings.replace(marker, b"no maximum lifetime", 1)})
    rewrite_manifest(root)

    verifier = root / "scripts/report_generation/verify_vibe_reports.py"
    text = verifier.read_text()
    old = re.search(r"workbook_semantic_sha\(wb\)=='([0-9a-f]{64})'", text)
    assert old
    verifier.write_text(text.replace(old.group(1), workbook_semantic_sha(workbook_path), 1))

    result = run_verifier(root)
    assert result.returncode != 0
    assert "preview JWT physical lifetime" in result.stdout
