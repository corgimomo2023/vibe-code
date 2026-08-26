import os
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from openpyxl import load_workbook


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
ARTIFACT_ROOT = Path(os.environ.get("VIBE_CODE_ROOT", REPOSITORY_ROOT))
REPORTS = ARTIFACT_ROOT / "wiki/product/reports"
PRD = REPORTS / "vibe-code-no-code-poc-prd-and-delivery-plan.docx"
TECH = REPORTS / "vibe-code-poc-technical-design-and-runtime-architecture.docx"
XLSX = REPORTS / "vibe-code-poc-data-specification-and-schema.xlsx"


def document_text(document: DocumentType) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def table_by_header(document: DocumentType, header: tuple[str, ...]):
    matches = []
    for table in document.tables:
        actual = tuple(cell.text.strip() for cell in table.rows[0].cells)
        if actual == header:
            matches.append(table)
    assert len(matches) == 1, (header, len(matches))
    return matches[0]


def headings(document: DocumentType) -> set[str]:
    return {
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and (paragraph.style.name or "").startswith("Heading")
    }


def numbered_ids(text: str, prefix: str) -> set[str]:
    return set(re.findall(rf"{prefix}-\d{{2}}", text))


def test_prd_contains_complete_use_cases_business_rules_and_visual_contracts():
    document = Document(PRD)
    document_headings = headings(document)

    assert "Vibe page layout" in document_headings
    assert "Business Rule Register" in document_headings
    assert "Authoritative state diagrams" in document_headings

    summary = table_by_header(document, ("ID", "Use Case", "Actors", "Stories"))
    assert len(summary.rows) == 10

    expected_fields = {
        "Actors",
        "Module",
        "Layer",
        "Preconditions",
        "Main flow",
        "Business Rules",
        "Security Guard",
        "Failure / Recovery",
        "Acceptance evidence",
        "Story trace",
    }
    detail_tables = []
    for table in document.tables:
        header = tuple(cell.text.strip() for cell in table.rows[0].cells)
        if header != ("Field", "Definition"):
            continue
        labels = {row.cells[0].text.strip() for row in table.rows[1:]}
        if "Main flow" in labels:
            detail_tables.append(table)
            assert labels == expected_fields
    assert len(detail_tables) == 9

    rules = table_by_header(
        document,
        ("ID", "Rule", "Applies to", "Enforcement / evidence"),
    )
    rule_ids = {row.cells[0].text.strip() for row in rules.rows[1:]}
    assert len(rule_ids) >= 15
    assert all(rule_id.startswith("BR-") for rule_id in rule_ids)

    captions = [paragraph.text for paragraph in document.paragraphs]
    assert sum(text.startswith("Figure UC-") for text in captions) == 9
    assert sum(text.startswith("Figure S-") for text in captions) == 10
    assert len(document.inline_shapes) >= 21


def test_technical_design_contains_state_transaction_and_ui_realization_diagrams():
    document = Document(TECH)
    document_headings = headings(document)

    assert "Vibe page component and layout contract" in document_headings
    assert "Authoritative state diagrams" in document_headings
    assert "Transactional sequence diagrams" in document_headings

    realization = table_by_header(
        document,
        (
            "Use Case",
            "Stories / Acceptance",
            "Business Rules",
            "Technical realization",
            "Verification",
        ),
    )
    assert len(realization.rows) == 10

    captions = [paragraph.text for paragraph in document.paragraphs]
    assert sum(text.startswith("Figure TS-") for text in captions) == 10
    assert sum(text.startswith("Figure TX-") for text in captions) >= 3
    assert len(document.inline_shapes) >= 16


def test_business_rule_mapping_is_exact_across_prd_register_and_technical_trace():
    prd = Document(str(PRD))
    technical = Document(str(TECH))
    summary = table_by_header(prd, ("ID", "Use Case", "Actors", "Stories"))
    uc_order = [row.cells[0].text.strip() for row in summary.rows[1:]]

    detail_tables = []
    for table in prd.tables:
        if tuple(cell.text.strip() for cell in table.rows[0].cells) != ("Field", "Definition"):
            continue
        values = {row.cells[0].text.strip(): row.cells[1].text for row in table.rows[1:]}
        if "Main flow" in values:
            detail_tables.append(values)
    assert len(detail_tables) == len(uc_order) == 9
    detail_mapping = {
        uc: numbered_ids(values["Business Rules"], "BR")
        for uc, values in zip(uc_order, detail_tables, strict=True)
    }

    register = table_by_header(prd, ("ID", "Rule", "Applies to", "Enforcement / evidence"))
    register_mapping = {uc: set() for uc in uc_order}
    for row in register.rows[1:]:
        rule_id = row.cells[0].text.strip()
        for uc in numbered_ids(row.cells[2].text, "UC"):
            register_mapping[uc].add(rule_id)

    trace = table_by_header(
        technical,
        ("Use Case", "Stories / Acceptance", "Business Rules", "Technical realization", "Verification"),
    )
    technical_mapping = {
        row.cells[0].text.strip(): numbered_ids(row.cells[2].text, "BR")
        for row in trace.rows[1:]
    }
    assert detail_mapping == register_mapping == technical_mapping


def test_platform_architecture_uses_nestjs_prisma_adminjs_and_redis_sessions():
    prd_text = document_text(Document(str(PRD)))
    technical = Document(str(TECH))
    technical_text = document_text(technical)
    combined = prd_text + "\n" + technical_text

    for required in (
        "NestJS",
        "Prisma",
        "AdminJS",
        "Redis",
        "opaque HttpOnly Secure session cookie",
        "server-side session",
        "short-lived JWT",
        "CSRF",
    ):
        assert required in technical_text

    for forbidden in (
        "Django", "DRF", "Django Admin", "django_session", "auth_user",
        "guest preview", "preview ticket", "one-time ticket",
    ):
        assert forbidden not in combined

    authentication = table_by_header(
        technical,
        ("Flow", "Credential", "Store / authority", "Lifetime", "Revocation / failure"),
    )
    assert len(authentication.rows) >= 6

    workbook = load_workbook(XLSX, read_only=True, data_only=False)
    assert "Redis Key Catalogue" in workbook.sheetnames
    redis_sheet = workbook["Redis Key Catalogue"]
    headers = tuple(redis_sheet.cell(6, column).value for column in range(1, 7))
    assert headers == (
        "Key Pattern",
        "Purpose",
        "Value Schema",
        "TTL / Bound",
        "Owner",
        "Invalidation / Failure",
    )
    rows = list(redis_sheet.iter_rows(min_row=7, values_only=True))
    assert len(rows) >= 4
    assert any(str(row[0]).startswith("session:") for row in rows)
    assert any("preview" in str(row[0]).lower() for row in rows)

    workbook_text = "\n".join(
        str(value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    for required in ("NestJS", "Prisma", "AdminJS", "Redis"):
        assert required in workbook_text
    for forbidden in ("Django", "DRF", "django_session", "auth_user"):
        assert forbidden not in workbook_text


def test_state_renderer_accepts_repository_contract(tmp_path: Path) -> None:
    import importlib.util
    import json

    renderer_path = REPOSITORY_ROOT / "scripts/report_generation/report_diagrams.py"
    spec = importlib.util.spec_from_file_location("report_diagrams_under_test", renderer_path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    contract_path = REPOSITORY_ROOT / "scripts/report_generation/state_contract.json"
    contract = json.loads(contract_path.read_text())
    preview_machine = next(
        machine for machine in contract["machines"]
        if machine["field"] == "PREVIEW_AUTH_GRANT.grant_status"
    )
    assert "JWT" in preview_machine["guard_summary"]
    assert "hashed jti" in preview_machine["guard_summary"]
    assert "ticket" not in preview_machine["guard_summary"].lower()

    output = tmp_path / "state.png"
    module.create_state_diagram(output, contract["machines"][0])
    assert output.is_file()
    assert output.stat().st_size > 10_000
